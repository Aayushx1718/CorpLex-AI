from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from app.extraction import DOCX_MIME, PDF_MIME, ExtractionError, extract, obligation_candidates
from generate_samples import as_docx


@pytest.mark.parametrize('fixture,mime,filename', [('docx_bytes', DOCX_MIME, 'vendor.docx'), ('pdf_bytes', PDF_MIME, 'vendor.pdf')])
def test_contract_fields_and_exact_evidence(request, fixture, mime, filename):
    result = extract(request.getfixturevalue(fixture), filename, mime)
    assert str(result.effective_date) == '2026-03-01'
    assert str(result.expiration_date) == '2027-02-28'
    assert [e['normalized_value'] for e in result.entities if e['kind'] == 'money'] == ['12000.00', '100.00', '1000.00']
    assert len([e for e in result.entities if e['kind'] == 'party']) == 2
    assert len(result.clauses) == 5  # preamble plus four numbered clauses
    for item in result.clauses:
        assert result.text[item['char_start']:item['char_end']] == item['text_content']
    for item in result.entities:
        assert result.text[item['char_start']:item['char_end']] == item['raw_text']
    delay = next(c for c in result.clauses if '3. Delay' in c['text_content'])
    assert obligation_candidates(delay, result.entities)[0]['amount'] is None


@pytest.mark.parametrize('data,name,mime', [
    (b'', 'empty.pdf', PDF_MIME), (b'not a pdf', 'bad.pdf', PDF_MIME),
    (b'PKbad zip', 'bad.docx', DOCX_MIME), (b'%PDF-', 'file.exe', PDF_MIME),
    (b'%PDF-', 'file.pdf', DOCX_MIME),
])
def test_rejects_invalid_files(data, name, mime):
    with pytest.raises(ExtractionError):
        extract(data, name, mime)


def test_rejects_encrypted_and_image_only_pdf(pdf_bytes):
    encrypted = BytesIO()
    writer = PdfWriter()
    writer.append(PdfReader(BytesIO(pdf_bytes)))
    writer.encrypt('test-password')
    writer.write(encrypted)
    with pytest.raises(ExtractionError, match='Encrypted'):
        extract(encrypted.getvalue(), 'locked.pdf', PDF_MIME)
    blank = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(blank)
    with pytest.raises(ExtractionError, match='OCR'):
        extract(blank.getvalue(), 'scan.pdf', PDF_MIME)


def test_rejects_expansion_bomb():
    output = BytesIO()
    with ZipFile(output, 'w', compression=ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
        archive.writestr('word/document.xml', 'x' * 1_000_000)
    with pytest.raises(ExtractionError, match='compressed'):
        extract(output.getvalue(), 'bomb.docx', DOCX_MIME)


@pytest.mark.parametrize('source', [
    'Effective date: 2026-02-30.',
    'Effective date: 2026-03-01. Expiration date: 2026-02-01.',
])
def test_rejects_invalid_dates(source):
    with pytest.raises(ExtractionError, match='date'):
        extract(as_docx(source), 'dates.docx', DOCX_MIME)


def test_docx_table_order_and_no_pagination_claim():
    document = Document()
    document.add_paragraph('1. Fees. Pay the following amount.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'USD 250.00'
    table.cell(0, 1).text = 'monthly'
    document.add_paragraph('2. Notice. Provide 30 days notice.')
    output = BytesIO()
    document.save(output)
    result = extract(output.getvalue(), 'table.docx', DOCX_MIME)
    assert [u['kind'] for u in result.units] == ['paragraph', 'table', 'paragraph']
    assert all(c['page_number'] is None for c in result.clauses)
    assert 'USD 250.00' in result.clauses[0]['text_content']


@pytest.mark.parametrize('name', ['nda', 'employment', 'vendor', 'service'])
@pytest.mark.parametrize('format', ['pdf', 'docx'])
def test_all_synthetic_contract_types_against_ground_truth(name, format):
    import json
    from pathlib import Path
    from generate_samples import as_pdf
    samples = Path(__file__).resolve().parents[2] / 'samples/contracts'
    expected = json.loads((samples / 'expected.json').read_text())[name]
    source = (samples / f'{name}.txt').read_text()
    data = as_pdf(source) if format == 'pdf' else as_docx(source)
    extracted = extract(data, f'{name}.{format}', PDF_MIME if format == 'pdf' else DOCX_MIME)
    assert extracted.effective_date.isoformat() == expected['effective_date']
    assert (extracted.expiration_date.isoformat() if extracted.expiration_date else None) == expected['expiration_date']
    assert [e['normalized_value'] for e in extracted.entities if e['kind'] == 'money'] == expected['money']
    assert len(extracted.clauses) == 5
    assert len([e for e in extracted.entities if e['kind'] == 'party']) == 2


def test_multipage_pdf_source_anchors():
    from generate_samples import as_pdf
    source = '1. First clause. ' + '\n'.join(['Example content.'] * 60) + '\n2. Payment. Pay USD 100.00.'
    result = extract(as_pdf(source), 'pages.pdf', PDF_MIME)
    assert len(result.units) == 2
    assert result.clauses[1]['page_number'] == 2
    for clause in result.clauses:
        assert result.text[clause['char_start']:clause['char_end']] == clause['text_content']
